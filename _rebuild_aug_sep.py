#!/usr/bin/env python3
"""Re-ingest Aug (3) + Sep (2), aliases/merges, NEWMAY, VERIFY."""
import re, json, os, unicodedata
from collections import defaultdict
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
HTML = os.path.join(HERE, "index.html")
AUG = "/Users/yuliia.nikolaieva/Downloads/Enhanced_Merged_Dataset_UA_Aug2026 (3).xlsx"
SEP = "/Users/yuliia.nikolaieva/Downloads/Enhanced_Merged_Dataset_UA_Sep2026 (2).xlsx"
JUNX = "/Users/yuliia.nikolaieva/Downloads/Enhanced_Merged_Dataset_UA_Jun2026.xlsx"

s = open(HTML, encoding="utf-8").read()

def gl(n):
    return json.loads(re.search(r"const " + n + r"=(\[.*?\]);", s, re.S).group(1))

def go(n):
    return json.loads(re.search(r"const " + n + r"=(\{.*?\});", s, re.S).group(1))

B, C, D = gl("B"), gl("C"), gl("D")
OV = go("OV")
VERIFY = go("VERIFY")

ROWLEN = 56
AUG_OFF = 44
SEP_OFF = 50

def norm(x):
    return re.sub(r"\s+", " ", str(x)).strip().upper()

def num(v):
    try:
        return float(str(v).replace(",", "").strip() or 0)
    except Exception:
        return 0.0

RENAME, VERT = {}, {}

def load_mistakes(path):
    if not os.path.exists(path):
        return
    wb = openpyxl.load_workbook(path, read_only=True)
    if "mistakes" in wb.sheetnames:
        for r in wb["mistakes"].iter_rows(values_only=True):
            r = list(r) + [""] * 12

            def sv(x):
                return "" if x is None else str(x).strip()

            asis, should, vto = sv(r[2]), sv(r[3]), sv(r[7])
            if asis in ("", "Merged Brand Name", "AS IS"):
                continue
            if should:
                RENAME[norm(asis)] = should
            if vto:
                VERT[norm(asis)] = vto

def load_comments(path):
    if not os.path.exists(path):
        return
    wb = openpyxl.load_workbook(path, read_only=True)
    if "comments" not in wb.sheetnames:
        return
    for r in wb["comments"].iter_rows(values_only=True):
        r = list(r) + [""] * 12

        def sv(x):
            if x is None:
                return ""
            t = str(x).strip()
            return "" if t.lower() in ("nan", "none") else t

        asis, should = sv(r[0]), sv(r[5])
        if asis in ("", "Merged Brand Name"):
            continue
        if should and norm(asis) != norm(should):
            RENAME[norm(asis)] = should
        if sv(r[7]):
            VERT[norm(should or asis)] = sv(r[7])

load_mistakes(JUNX)
load_comments(AUG)
load_comments(SEP)

ALIAS_RAW = {
    "Фора Преміум": "Fora",
    "Fora / ФОРА": "Fora",
    "АНЦ": "ANC",
    "Аптека АНЦ": "ANC",
    "АПТЕКА АНЦ": "ANC",
    "ATB Black": "ATB",
    "ATB / АТБ": "ATB",
    "АТБ": "ATB",
    "Blyzenko / Близенько": "Blyzenko",
    "Біла Ромашка": "Bila Romashka",
    "Аптека Біла Ромашка": "Bila Romashka",
    "Bila Romashka / Біла Ромашка": "Bila Romashka",
    "Аптека Подорожник": "Подорожник",
    "EKO market/ЕКО маркет": "EKO market",
    "LOKO від Сільпо": "LOKO",
    "Loko": "LOKO",
    "Kopiyka": "KOPIYKA",
    "Kopiyka / Копійка": "KOPIYKA",
    "Hop Hey": "HOP HEY",
    "Hop Hey MAX": "HOP HEY",
    "Beer Market": "BEER MARKET",
    "Rukavychka": "RUKAVYCHKA",
    "Pyvna Boroda": "PYVNA BORODA",
    "Vapors": "VAPORS",
    "АПТЕКА 9-1-1": "APTEKA 911",
    "Аптека D.S.": "APTEKA D.S",
    "Аптека Доброго Дня": "Apteka DD",
    "Apteka Dobroho Dnia": "Apteka DD",
    "PYVNE REMESLO": "REMESLO BREWERY",
    "LEPRUKON": "LEPRUKON",
    "Leprekon": "LEPRUKON",
    "Копійка від АНЦ": "Kopiyka by ANC",
    "Благодія": "Blagodiya by ANC",
    "Шара від АНЦ": "Shara by ANC",
    "VAPE SHOP KYIV": "VAPERY | VAPE SHOP",
    "Solo Vino / Соло Вино": "SOLO VINO",
    "UZBETSKYI PLOV №1": "UZBETSKYI PLOV #1 / УЗБЕЦЬКИЙ ПЛОВ #1",
    "UZBETSKYI PLOV": "UZBETSKYI PLOV #1 / УЗБЕЦЬКИЙ ПЛОВ #1",
    "MUZA BY KLOPOTENKO": "МУЗА ВІД ЄВГЕНА КЛОПОТЕНКА / MUZA VID YEVHENA KLOPOTENKA",
    "MUZA BY KLOPOTENKO LVIV": "МУЗА ВІД ЄВГЕНА КЛОПОТЕНКА / MUZA VID YEVHENA KLOPOTENKA",
    "beerstore": "BEERSTORE",
    "Beerstore": "BEERSTORE",
    "BRSM": "Good market - БРСМ",
    "Галицький Кухоль": "HALYTSKYI KUKHOL",
    "Дім Свят LeBouquet": "LeBouquet",
    "LE BOUQUET": "LeBouquet",
    "Jerkins та Пиво": "JERKINS AND BEER",
    "Пара Рівне": "PARA RIVNE",
    "ПАРА РІВНЕ": "PARA RIVNE",
    "PARA RIVNE": "PARA RIVNE",
    "Ч Магазин Крафтових Напоїв": "КВАЧ Магазин Крафтових Напоїв",
    "PINK FLOWER": "PINK FLOWER",
    "ПАРАЛАБ": "ПАРАЛАБ",
    "KHMELNYI LEV": "KHMELNYI LEV",
}

CANON = {}
for v in list(RENAME.values()) + list(ALIAS_RAW.values()):
    CANON[norm(v)] = v
ALIAS = {norm(k): norm(v) for k, v in ALIAS_RAW.items()}
for k, v in RENAME.items():
    ALIAS[k] = norm(v)

def canon(name):
    n = norm(name)
    t = ALIAS.get(n) or CANON.get(n) or name
    n2 = norm(t)
    return ALIAS.get(n2) or CANON.get(n2) or t

CITY_ALIAS = {"Zaporizhzhia": "Zaporizhia", "Irpin": "Irpin-Bucha"}

def acity(x):
    x = "" if x is None else str(x).strip()
    return CITY_ALIAS.get(x, x)

cyr = {
    "А": "A", "Б": "B", "В": "V", "Г": "H", "Ґ": "G", "Д": "D", "Е": "E", "Є": "E",
    "Ж": "ZH", "З": "Z", "И": "Y", "І": "I", "Ї": "I", "Й": "Y", "К": "K", "Л": "L",
    "М": "M", "Н": "N", "О": "O", "П": "P", "Р": "R", "С": "S", "Т": "T", "У": "U",
    "Ф": "F", "Х": "H", "Ц": "C", "Ч": "CH", "Ш": "SH", "Щ": "SCH", "Ь": "", "Ю": "U",
    "Я": "A", "Ъ": "", "Ы": "Y", "Э": "E",
}

def dkey(x):
    x = unicodedata.normalize("NFKD", str(x))
    x = "".join(c for c in x if not unicodedata.combining(c))
    x = x.replace("№", "#").replace("Nº", "#")
    return re.sub(r"[^A-Za-z0-9А-Яа-яЇїІіЄєҐґ#]", "", x).upper()

def tkey(x):
    x = unicodedata.normalize("NFKD", str(x))
    x = "".join(c for c in x if not unicodedata.combining(c)).upper()
    x = x.replace("№", "#").replace("Nº", "#")
    x = "".join(cyr.get(ch, ch) for ch in x)
    x = re.sub(r"[^A-Z0-9]", "", x)
    return (
        x.replace("Z", "S")
        .replace("KH", "H")
        .replace("IY", "I")
        .replace("YI", "I")
        .replace("W", "V")
        .replace("IA", "A")
        .replace("YA", "A")
        .replace("EE", "E")
    )

live = [i for i, b in enumerate(B) if b[0] != "__merged__"]
by_norm, by_dkey, by_tkey = {}, {}, {}
for i in live:
    by_norm.setdefault(norm(canon(B[i][0])), i)
    by_dkey.setdefault(dkey(B[i][0]), i)
    tk = tkey(B[i][0])
    if len(tk) >= 6:
        by_tkey.setdefault(tk, i)

def resolve_bi(raw_display):
    cn = canon(raw_display)
    n = norm(cn)
    if n in by_norm:
        return by_norm[n]
    dk = dkey(cn)
    if dk in by_dkey:
        return by_dkey[dk]
    tk = tkey(cn)
    if len(tk) >= 6 and tk in by_tkey:
        return by_tkey[tk]
    return None

FOOD_CATS = {
    "Restaurant", "Fast Food", "Pizza", "Burgers", "Sushi", "Shawarma/Kebab", "Coffee & Bakery",
}
RETAIL_CATS = {
    "Grocery", "Other Retail", "Pharma", "Alcohol", "Tobacco/Vape", "Flowers & Gifts", "Pets",
    "Beauty", "Dark Store", "Meat", "Specialty", "Toys", "Erotic", "Gas Station",
}

for asisnorm, v in VERT.items():
    bi = resolve_bi(asisnorm)
    if bi is None:
        continue
    vv = v.lower()
    if vv == "food":
        B[bi][1] = "r"
        if len(B[bi]) > 2 and B[bi][2] in RETAIL_CATS:
            B[bi][2] = "Restaurant"
    elif vv.startswith("store"):
        B[bi][1] = "t"
        if len(B[bi]) > 2 and B[bi][2] in FOOD_CATS:
            B[bi][2] = "Other Retail"

# Normalize row length; clear Aug+Sep slots
for d in D:
    if len(d) < ROWLEN:
        d.extend([0] * (ROWLEN - len(d)))
    elif len(d) > ROWLEN:
        d[:] = d[:ROWLEN]
    d[AUG_OFF : SEP_OFF + 6] = [0] * 12

city_idx = {c: i for i, c in enumerate(C)}
row_idx = {(d[0], d[1]): i for i, d in enumerate(D)}
existing_norm = {norm(B[i][0]): i for i in live}

def ingest_workbook(path, col_off, track_seen=None):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["Data"]
    it = ws.iter_rows(values_only=True)
    hdr = [str(c).strip() if c else "" for c in next(it)]
    ix = {h: i for i, h in enumerate(hdr)}
    ORD = defaultdict(lambda: {"t": 0.0, "g": 0.0, "b": 0.0, "gx": 0, "both": 0, "bx": 0})
    disp_new, vert_new = {}, {}
    blank_t = 0.0
    brand_tot = defaultdict(float)

    for r in it:
        raw = r[ix["Merged Brand Name"]]
        if raw is None or str(raw).strip() == "":
            blank_t += num(r[ix["Total orders"]])
            continue
        raw = str(raw).strip()
        cn = canon(raw)
        nb = norm(cn)
        if track_seen is not None:
            track_seen.add(nb)
            brand_tot[nb] += num(r[ix["Total orders"]])
        city = acity(r[ix["Merged city"]])
        bi = resolve_bi(raw)
        if bi is not None:
            key = ("bi", bi, city)
        else:
            key = ("new", nb, city)
            disp_new.setdefault(nb, cn)
            vert_new.setdefault(nb, str(r[ix["Delivery Vertical"]] or ""))
        g = num(r[ix["Glovo orders"]])
        b = num(r[ix["Bolt Orders"]])
        t = num(r[ix["Total orders"]])
        o = str(r[ix["Overlap of exclusive"]] or "").strip().lower()
        a = ORD[key]
        a["t"] += t
        a["g"] += g
        a["b"] += b
        if "glovo" in o and "exclusive" in o:
            if g > 0:
                a["gx"] += 1
        elif "bolt" in o and "exclusive" in o:
            if b > 0:
                a["bx"] += 1
        elif "overlap" in o or o in ("g + b", "g+b"):
            if t > 0:
                a["both"] += 1

    newb = 0
    for key, a in ORD.items():
        kind, id_, city = key
        if city not in city_idx:
            city_idx[city] = len(C)
            C.append(city)
        ci = city_idx[city]
        if kind == "bi":
            bi = id_
        else:
            nb = id_
            if nb in existing_norm:
                bi = existing_norm[nb]
            else:
                vv = vert_new.get(nb, "")
                tp = "r" if vv.lower().startswith("food") else "t"
                cat = "Restaurant" if tp == "r" else "Other Retail"
                bi = len(B)
                B.append([disp_new[nb], tp, cat])
                existing_norm[nb] = bi
                by_norm[nb] = bi
                by_dkey[dkey(disp_new[nb])] = bi
                tk = tkey(disp_new[nb])
                if len(tk) >= 6:
                    by_tkey[tk] = bi
                live.append(bi)
                newb += 1
        vals = [round(a["t"]), round(a["g"]), round(a["b"]), a["gx"], a["both"], a["bx"]]
        ri = row_idx.get((bi, ci))
        if ri is None:
            D.append([bi, ci] + [0] * (ROWLEN - 2))
            ri = len(D) - 1
            row_idx[(bi, ci)] = ri
        D[ri][col_off : col_off + 6] = vals

    tot = sum(d[col_off] for d in D)
    print(path.split("/")[-1], "col", col_off, "total", int(tot), "new brands", newb, "blank", round(blank_t))
    return tot, brand_tot

aug_seen = set()
aug_tot, aug_brand_tot = ingest_workbook(AUG, AUG_OFF, aug_seen)
sep_tot, _ = ingest_workbook(SEP, SEP_OFF, None)

# Manual type fixes (by display name norm)
TYPE_FIX = {
    "СПІЛЬНЕ.КАВА": ("r", "Coffee & Bakery"),
    "БУФЕТ": ("r", "Restaurant"),
}
for i, b in enumerate(B):
    if b[0] == "__merged__":
        continue
    fx = TYPE_FIX.get(norm(b[0]))
    if fx:
        b[1], b[2] = fx

def bi_by_name(name):
    n = norm(name)
    if n in by_norm:
        return by_norm[n]
    for i, b in enumerate(B):
        if b[0] == "__merged__":
            continue
        if norm(b[0]) == n:
            return i
    return None

def merge_brand(from_name, to_name):
    fi = bi_by_name(from_name)
    ti = bi_by_name(to_name)
    if fi is None or ti is None or fi == ti:
        print("merge skip", from_name, to_name, fi, ti)
        return
    drop = []
    merged = 0
    for i, d in enumerate(D):
        if d[0] != fi:
            continue
        ci = d[1]
        ri = row_idx.get((ti, ci))
        if ri is None:
            d[0] = ti
            row_idx[(ti, ci)] = i
        else:
            for c in range(2, ROWLEN):
                D[ri][c] += d[c]
            drop.append(i)
        merged += 1
    for i in sorted(drop, reverse=True):
        D.pop(i)
    row_idx.clear()
    for i, d in enumerate(D):
        row_idx[(d[0], d[1])] = i
    B[fi][0] = "__merged__"
    print("merged", B[ti][0], "<-", from_name, "rows", merged)

merge_brand("Jerkins та Пиво", "JERKINS AND BEER")
merge_brand("Дім Свят LeBouquet", "LE BOUQUET")
lb = bi_by_name("LE BOUQUET")
if lb is not None:
    B[lb][0] = "LeBouquet"
    by_norm[norm("LeBouquet")] = lb

aug_zero_seen = {nb for nb, t in aug_brand_tot.items() if t == 0}

def btype_for_raw(raw):
    bi = resolve_bi(raw)
    if bi is not None:
        return B[bi][1]
    nb = norm(canon(raw))
    i = existing_norm.get(nb)
    return B[i][1] if i is not None else "t"

# OV September
S = defaultdict(lambda: {"ges": 0, "gbs": 0, "bes": 0, "geo": 0.0, "gbo": 0.0, "beo": 0.0})
wb = openpyxl.load_workbook(SEP, read_only=True)
ws = wb["Data"]
it = ws.iter_rows(values_only=True)
hdr = [str(c).strip() if c else "" for c in next(it)]
ix = {h: i for i, h in enumerate(hdr)}
for r in it:
    raw = r[ix["Merged Brand Name"]]
    if raw is None or str(raw).strip() == "":
        continue
    raw = str(raw).strip()
    tp = btype_for_raw(raw)
    o = str(r[ix["Overlap of exclusive"]] or "").strip().lower()
    g = num(r[ix["Glovo orders"]])
    b = num(r[ix["Bolt Orders"]])
    t = num(r[ix["Total orders"]])
    cls = 0 if ("glovo" in o and "exclusive" in o) else (2 if ("bolt" in o and "exclusive" in o) else 1)
    for sg in ("all", "restaurant", "retail"):
        if sg == "all" or (sg == "restaurant" and tp == "r") or (sg == "retail" and tp == "t"):
            d = S[sg]
            if cls == 0:
                d["geo"] += t
                if g > 0:
                    d["ges"] += 1
            elif cls == 2:
                d["beo"] += t
                if b > 0:
                    d["bes"] += 1
            else:
                d["gbo"] += t
                if t > 0:
                    d["gbs"] += 1
OV["sp"] = {
    sg: {
        "ges": S[sg]["ges"],
        "gbs": S[sg]["gbs"],
        "bes": S[sg]["bes"],
        "geo": round(S[sg]["geo"]),
        "gbo": round(S[sg]["gbo"]),
        "beo": round(S[sg]["beo"]),
    }
    for sg in ("all", "restaurant", "retail")
}

aug_b, aug_bc = set(), set()
for d in D:
    if B[d[0]][0] == "__merged__":
        continue
    if d[AUG_OFF] > 0 or (d[AUG_OFF + 3] + d[AUG_OFF + 4] + d[AUG_OFF + 5]) > 0:
        nb = norm(B[d[0]][0])
        aug_b.add(nb)
        aug_bc.add((nb, C[d[1]]))

jb, jbc, jd = defaultdict(float), defaultdict(float), {}
for d in D:
    if B[d[0]][0] == "__merged__":
        continue
    if d[SEP_OFF] > 0 or (d[SEP_OFF + 3] + d[SEP_OFF + 4] + d[SEP_OFF + 5]) > 0:
        nb = norm(B[d[0]][0])
        jd.setdefault(nb, B[d[0]][0])
        jb[nb] += d[SEP_OFF]
        jbc[(nb, C[d[1]])] += d[SEP_OFF]

def junk(n):
    return "TEST" in n.upper() or n == "__MERGED__"

def is_new_partner(nb):
    if nb in aug_b:
        return False
    if nb in aug_seen:
        return False
    return True

npp = sorted(
    [(jd[b], round(o)) for b, o in jb.items() if is_new_partner(b) and not junk(b)],
    key=lambda x: -x[1],
)
nll = sorted(
    [
        (jd[b], c, round(o))
        for (b, c), o in jbc.items()
        if (b, c) not in aug_bc and is_new_partner(b) and not junk(b)
    ],
    key=lambda x: -x[2],
)

NEWMAY = {
    "newPartnerCount": len(npp),
    "newPartnerOrders": round(sum(p[1] for p in npp)),
    "topNewPartners": npp[:20],
    "newLocCount": len(nll),
    "newLocOrders": round(sum(l[2] for l in nll)),
    "topNewLocs": [[a, b, c] for a, b, c in nll[:20]],
    "newBrandCount": 0,
    "newCities": [],
    "vsMonth": "August",
    "curMonth": "September",
}

VERIFY["augTot"] = int(aug_tot)
VERIFY["augustSource"] = "Enhanced_Merged_Dataset_UA_Aug2026 (3).xlsx"
VERIFY["sepTot"] = int(sep_tot)
VERIFY["septemberSource"] = "Enhanced_Merged_Dataset_UA_Sep2026 (2).xlsx"
VERIFY["augSeenPartners"] = sorted(aug_seen)
VERIFY["mistakesApplied"] = True

print("NEW partners", len(npp), "new locs", len(nll))

def dump(x):
    return json.dumps(x, separators=(",", ":"), ensure_ascii=False)

def repl(t, n, cl, lit):
    pat = r"const " + n + (r"=\[.*?\];" if cl == "]" else r"=\{.*?\};")
    m = re.search(pat, t, re.S)
    return t[: m.start()] + lit + t[m.end() :]

s = repl(s, "B", "]", "const B=" + dump(B) + ";")
s = repl(s, "C", "]", "const C=" + dump(C) + ";")
s = repl(s, "D", "]", "const D=" + dump(D) + ";")
s = repl(s, "OV", "}", "const OV=" + dump(OV) + ";")
s = repl(s, "NEWMAY", "}", "const NEWMAY=" + dump(NEWMAY) + ";")
s = repl(s, "VERIFY", "}", "const VERIFY=" + dump(VERIFY) + ";")
open(HTML, "w", encoding="utf-8").write(s)
print("Wrote", HTML)
